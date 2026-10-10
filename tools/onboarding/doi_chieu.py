"""Bước E — đối chiếu công / lương khi chạy song song hệ thống cũ và BossKin (pilot).

So từng nhân viên (theo mã NV) giữa file xuất của BossKin và file của hệ thống cũ, ghi ra file Excel các dòng lệch.

    # Lương: file BossKin = Excel → Xuất → Bảng lương (tháng cần so, tiếng Việt)
    python3 tools/onboarding/doi_chieu.py luong --bosskin bang-luong.xlsx --cu luong-cu.xlsx --ra doi-chieu-luong.xlsx

    # Công: file BossKin = Excel → Xuất → Chấm công (cả kỳ), cộng theo nhân viên
    python3 tools/onboarding/doi_chieu.py cong --bosskin cham-cong.xlsx --cu cong-cu.xlsx --ra doi-chieu-cong.xlsx

File cũ (.xlsx hoặc .csv UTF-8) cần một dòng tiêu đề và một dòng mỗi nhân viên. Tên cột khác BossKin thì khai báo
bằng --cot "<cột BossKin>=<cột file cũ>" (lặp lại được), vd. --cot "Mã NV=Mã nhân viên" --cot "Net (thực lãnh)=Thực lĩnh".
Cột không có trong file cũ thì bỏ qua. Sai số cho phép: --sai-so (công/giờ, mặc định 0.01), --sai-so-tien (mặc định 1000 đ).

Chỉ cần thư viện openpyxl. Hướng dẫn: docs/chuyen-du-lieu.md mục "Chạy song song".
"""
import argparse
import csv
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Font, PatternFill

KEY = "Mã NV"

# Cột so sánh mặc định (nhãn tiếng Việt trong file xuất BossKin) và loại sai số
PAYROLL_COLS = OrderedDict([
    ("Ngày công", "num"), ("Giờ ngày", "num"), ("Giờ đêm", "num"), ("Giờ OT ngày", "num"), ("Giờ OT đêm", "num"),
    ("Lương cơ bản", "money"), ("Tổng trợ cấp", "money"), ("Tổng thu nhập", "money"), ("Gross", "money"),
    ("Tổng khấu trừ", "money"), ("Net (thực lãnh)", "money"),
])
# Công: cộng các dòng chấm công theo nhân viên
ATT_COLS = OrderedDict([("Ngày có công", "num"), ("Giờ thực tế", "num"), ("OT ngày", "num"), ("OT đêm", "num")])


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if s in ("", "-"):
        return None
    # 1.234.567 hoặc 1,234,567 hoặc 7,5 → số
    if s.count(",") == 1 and s.count(".") == 0 and len(s.split(",")[1]) <= 2:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3 and s.count(",") == 0):
        s = s.replace(".", "")
    s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def read_table(path):
    """Trả (headers, rows) của sheet đầu tiên có cột chứa 'Mã NV' (hoặc sheet đầu)."""
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            data = list(csv.reader(f))
        return data[0], [dict(zip(data[0], r)) for r in data[1:] if any(c.strip() for c in r)]
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    chosen = None
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if any(isinstance(h, str) and h.strip().startswith(KEY) for h in first):
            chosen = ws
            break
    ws = chosen or wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = [dict(zip(headers, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    return headers, out


def find_col(headers, name):
    """Khớp tên cột, bỏ qua dấu (*) và khoảng trắng thừa."""
    norm = lambda h: h.replace("(*)", "").replace("*", "").strip().lower()
    for h in headers:
        if norm(h) == norm(name):
            return h
    return None


def index_by_code(headers, rows, key_name, label):
    key = find_col(headers, key_name)
    if key is None:
        sys.exit(f"{label}: không thấy cột mã nhân viên '{key_name}' (khai báo bằng --cot \"{KEY}=<tên cột>\")")
    out = OrderedDict()
    for r in rows:
        code = r.get(key)
        if code is None or str(code).strip() == "":
            continue
        out.setdefault(str(code).strip(), []).append(r)
    return out


def aggregate_attendance(headers, rows):
    """Cộng file chấm công BossKin theo nhân viên → (headers, rows) dạng tổng hợp."""
    by = index_by_code(headers, rows, KEY, "BossKin")
    h_in, h_hours = find_col(headers, "Check-in"), find_col(headers, "Giờ thực tế")
    h_ot, h_otn = find_col(headers, "OT ngày"), find_col(headers, "OT đêm")
    h_name = find_col(headers, "Họ và tên")
    out = []
    for code, rs in by.items():
        days = {r.get(find_col(headers, "Ngày")) for r in rs if h_in and r.get(h_in) not in (None, "")}
        out.append({
            KEY: code,
            "Họ và tên": rs[0].get(h_name) if h_name else "",
            "Ngày có công": float(len(days)),
            "Giờ thực tế": sum(to_num(r.get(h_hours)) or 0 for r in rs) if h_hours else None,
            "OT ngày": sum(to_num(r.get(h_ot)) or 0 for r in rs) if h_ot else None,
            "OT đêm": sum(to_num(r.get(h_otn)) or 0 for r in rs) if h_otn else None,
        })
    return [KEY, "Họ và tên", *ATT_COLS.keys()], out


def compare(mode, bk_path, old_path, mapping, tol_num, tol_money):
    bk_headers, bk_rows = read_table(bk_path)
    cols = PAYROLL_COLS
    if mode == "cong":
        bk_headers, bk_rows = aggregate_attendance(bk_headers, bk_rows)
        cols = ATT_COLS
    old_headers, old_rows = read_table(old_path)

    bk = index_by_code(bk_headers, bk_rows, KEY, "BossKin")
    old = index_by_code(old_headers, old_rows, mapping.get(KEY, KEY), "File cũ")
    name_col = find_col(bk_headers, "Họ và tên")

    compared = []   # (cột BossKin, cột cũ, loại)
    for col, kind in cols.items():
        b = find_col(bk_headers, col)
        o = find_col(old_headers, mapping.get(col, col))
        if b and o:
            compared.append((b, o, kind))
    if not compared:
        sys.exit("Không có cột nào khớp giữa hai file — khai báo --cot \"<cột BossKin>=<cột file cũ>\"")

    diffs, matched = [], 0
    for code, rs in bk.items():
        if code not in old:
            continue
        b_row, o_rows = rs[0], old[code]
        row_ok = True
        for b, o, kind in compared:
            bv = to_num(b_row.get(b))
            ov = sum(to_num(r.get(o)) or 0 for r in o_rows) if any(to_num(r.get(o)) is not None for r in o_rows) else None
            tol = tol_money if kind == "money" else tol_num
            if (bv or 0) == (ov or 0) or (bv is not None and ov is not None and abs(bv - ov) <= tol):
                continue
            if bv is None and ov is None:
                continue
            row_ok = False
            diffs.append([code, b_row.get(name_col, "") if name_col else "", b, bv, ov,
                          None if bv is None or ov is None else round(bv - ov, 2)])
        matched += row_ok
    only_bk = [[c, rs[0].get(name_col, "") if name_col else ""] for c, rs in bk.items() if c not in old]
    only_old = [[c] for c in old if c not in bk]
    return {"compared": compared, "diffs": diffs, "matched": matched, "only_bk": only_bk, "only_old": only_old,
            "total_bk": len(bk), "total_old": len(old)}


def write_report(res, out_path, mode):
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DBEAFE")

    def sheet(title, headers, rows, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, len(str(h)) + 4)
        ws.freeze_panes = "A2"

    diff_people = len({d[0] for d in res["diffs"]})
    sheet("Tổng hợp", ["Mục", "Giá trị"], [
        ["Loại đối chiếu", "Lương" if mode == "luong" else "Công"],
        ["Nhân viên ở BossKin", res["total_bk"]],
        ["Nhân viên ở file cũ", res["total_old"]],
        ["Khớp hoàn toàn", res["matched"]],
        ["Có chênh lệch", diff_people],
        ["Chỉ có ở BossKin", len(res["only_bk"])],
        ["Chỉ có ở file cũ", len(res["only_old"])],
        ["Cột đã so", ", ".join(f"{b} ↔ {o}" for b, o, _ in res["compared"])],
    ], first=True)
    sheet("Chênh lệch", ["Mã NV", "Họ và tên", "Cột", "BossKin", "File cũ", "Lệch (BossKin − cũ)"], res["diffs"])
    sheet("Chỉ có ở BossKin", ["Mã NV", "Họ và tên"], res["only_bk"])
    sheet("Chỉ có ở file cũ", ["Mã NV"], res["only_old"])
    wb.save(out_path)
    return diff_people


def main(argv=None):
    p = argparse.ArgumentParser(description="Đối chiếu công / lương BossKin với hệ thống cũ")
    p.add_argument("loai", choices=["luong", "cong"])
    p.add_argument("--bosskin", required=True, help="File xuất từ BossKin (.xlsx)")
    p.add_argument("--cu", required=True, help="File của hệ thống cũ (.xlsx / .csv)")
    p.add_argument("--ra", required=True, help="File kết quả (.xlsx)")
    p.add_argument("--cot", action="append", default=[], help='"<cột BossKin>=<cột file cũ>"')
    p.add_argument("--sai-so", type=float, default=0.01)
    p.add_argument("--sai-so-tien", type=float, default=1000)
    a = p.parse_args(argv)
    mapping = {}
    for m in a.cot:
        if "=" not in m:
            p.error(f"--cot sai dạng: {m}")
        k, v = m.split("=", 1)
        mapping[k.strip()] = v.strip()
    res = compare(a.loai, a.bosskin, a.cu, mapping, a.sai_so, a.sai_so_tien)
    n = write_report(res, a.ra, a.loai)
    print(f"Đã so {len(res['compared'])} cột cho {res['total_bk']} NV BossKin / {res['total_old']} NV file cũ: "
          f"{res['matched']} khớp, {n} lệch, {len(res['only_bk'])} chỉ có ở BossKin, {len(res['only_old'])} chỉ có ở file cũ "
          f"→ {a.ra}")
    return 0 if n == 0 and not res["only_bk"] and not res["only_old"] else 1


if __name__ == "__main__":
    sys.exit(main())
