"""Bước E — tạo mẫu Excel nhập dữ liệu thật cho công ty mới.

Đầu vào là file xuất từ BossKin của một công ty trống (Excel → Xuất, chọn 8 loại: Phòng ban, Chức vụ, Bậc lương,
Danh sách ca, Loại nghỉ phép, Nhân viên (mọi cột nhập được), Số dư phép, Phân công ca). Nhờ vậy tên sheet / tên cột
luôn khớp đúng bản đang chạy. Script tách thành 4 file theo thứ tự nhập (bước sau dùng mã tạo ở bước trước, nên
màn xem trước chỉ đúng khi bước trước đã nhập xong), thêm sheet "Hướng dẫn" và điền dữ liệu ví dụ hư cấu vào bản vi-du/.

    python3 tools/onboarding/make_template.py <file-xuat.xlsx> <thư-mục-ra>

Hướng dẫn đầy đủ: docs/chuyen-du-lieu.md.
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

GUIDE = [
    ("Mẫu nhập dữ liệu BossKin", None),
    ("", None),
    ("Thứ tự nhập", "Nhập lần lượt 4 file: 1-phong-ban → 2-danh-muc (chức vụ, bậc lương, ca, loại phép) → 3-nhan-vien → "
                    "4-so-du-phep-phan-ca. Mỗi file nhập xong (không còn lỗi) mới sang file sau, vì file sau dùng mã của file trước."),
    ("Cột Lệnh", "MERGE = thêm mới hoặc cập nhật theo mã (nên dùng). CREATE chỉ thêm, UPDATE chỉ sửa, DELETE xoá (xoá mềm). "
                 "Dòng để trống cột Lệnh được hiểu là MERGE."),
    ("Cột có (*)", "Bắt buộc. Cột \"Tên …\" cạnh cột mã chỉ để đọc, hệ thống bỏ qua khi nhập."),
    ("Ngày", "Dạng dd/MM/yyyy, ví dụ 01/11/2026. Giờ dạng HH:mm, ví dụ 08:00."),
    ("Giới tính / trạng thái", "Giới tính: MALE, FEMALE, OTHER. Trạng thái NV: ACTIVE, INACTIVE, ON_LEAVE."),
    ("Tài khoản nhân viên", "Mỗi dòng Nhân viên có Email sẽ tạo tài khoản đăng nhập bằng email đó với mật khẩu ngẫu nhiên. "
                            "Nhập xong: Phân quyền → Người dùng → Xuất mật khẩu tạm để lấy file mật khẩu phát cho từng người; "
                            "mỗi người phải đổi mật khẩu ở lần đăng nhập đầu."),
    ("Lương hợp đồng", "Mức lương thoả thuận của từng người (để trống thì dùng mức của bậc lương)."),
    ("Số dư phép", "Phép còn lại lúc chuyển sang BossKin: ghi Số ngày phép = số ngày còn được nghỉ, Ngày cấp = ngày bắt đầu dùng BossKin, "
                   "Hiệu lực đến = hạn dùng (vd. 31/12/2026). Loại phép phải có ở sheet Loại nghỉ phép hoặc đã có trong hệ thống."),
    ("Phân công ca", "Mỗi dòng = một người, một ngày, một ca. Muốn xếp nhiều ngày nhanh hơn, dùng màn Ca làm (xoay ca / bảng ngang) sau khi nhập."),
    ("Ca linh hoạt", "Sheet Danh sách ca chỉ nhập ca cố định. Ca linh hoạt (đủ số giờ trong khung) tạo ở màn Ca làm."),
    ("Cách nhập", "Màn Excel (hoặc nút Nhập ở từng màn) → chọn file → xem trước (Tạo mới / Cập nhật / Không đổi / cảnh báo) → Xác nhận. "
                  "Tải file kết quả để xem dòng lỗi (cột Trạng thái / Thông báo), sửa rồi nhập lại: MERGE không tạo trùng."),
    ("Không nhập ở đây", "Chấm công cũ và bảng lương cũ: giữ ở hệ thống cũ, đối chiếu bằng tools/onboarding/doi_chieu.py khi chạy song song."),
]

EXAMPLE = {
    "Phòng ban": [
        ["MERGE", "BGD", "Ban giám đốc", None, None, None, None, 100, None],
        ["MERGE", "HCNS", "Hành chính nhân sự", "BGD", None, None, None, 100, None],
        ["MERGE", "XUONG", "Xưởng sản xuất", "BGD", "Xưởng chính", 10.7769, 106.7009, 200, None],
    ],
    "Chức vụ": [
        ["MERGE", "GD", "Giám đốc", "BGD", None, None],
        ["MERGE", "NVNS", "Nhân viên nhân sự", "HCNS", None, None],
        ["MERGE", "CN", "Công nhân", "XUONG", None, None],
    ],
    "Bậc lương": [
        ["MERGE", "VP1", "Văn phòng bậc 1", "DAILY", 9000000, None, None],
        ["MERGE", "CN1", "Công nhân bậc 1", "DAILY", 7000000, "XUONG", None],
    ],
    "Danh sách ca": [
        ["MERGE", "HC", "Hành chính", "08:00", "17:00", "12:00", "13:00", None, "#1976D2", None],
        ["MERGE", "SANG", "Ca sáng", "06:00", "14:00", "10:00", "10:30", "Xưởng", "#43A047", "XUONG"],
    ],
    "Loại nghỉ phép": [
        ["MERGE", "PN", "Phép năm", "DAY", "true", None, "true", "12 ngày / năm"],
        ["MERGE", "KL", "Nghỉ không lương", "DAY", "false", None, "true", None],
    ],
    "Nhân viên": [
        ["MERGE", "NV001", "Trần Minh Khoa", "khoa@vidu.test", "0901000001", "MALE", "15/04/1985", "01/03/2020", "ACTIVE",
         "BGD", None, "GD", None, "VP1", 30000000, "001085000001", "Q1, TP.HCM", None, "0011000000001", "Vietcombank",
         "8000000001", "7900000001", "BV Chợ Rẫy", "01/03/2020"],
        ["MERGE", "NV002", "Lê Thu Hà", "ha@vidu.test", "0901000002", "FEMALE", "02/09/1992", "10/06/2022", "ACTIVE",
         "HCNS", None, "NVNS", None, "VP1", 12000000, "001092000002", None, None, "0011000000002", "Vietcombank",
         None, "7900000002", None, "10/06/2022"],
        ["MERGE", "NV003", "Phạm Văn Lực", "luc@vidu.test", None, "MALE", None, "05/01/2024", "ACTIVE",
         "XUONG", None, "CN", None, "CN1", None, None, None, None, None, None, None, None, None, None],
    ],
    "Số dư phép": [
        ["MERGE", "NV001", None, None, "PN", None, 8, None, None, None, "01/11/2026", "01/11/2026", "31/12/2026", "Phép năm còn lại khi chuyển hệ thống"],
        ["MERGE", "NV002", None, None, "PN", None, 4.5, None, None, None, "01/11/2026", "01/11/2026", "31/12/2026", "Phép năm còn lại khi chuyển hệ thống"],
    ],
    "Phân công ca": [
        ["MERGE", None, "NV002", None, None, "02/11/2026", "HC", None, None, None, None, None, None],
        ["MERGE", None, "NV003", None, None, "02/11/2026", "SANG", None, None, None, None, None, None],
    ],
}


def add_guide(wb):
    ws = wb.create_sheet("Hướng dẫn", 0)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    for i, (k, v) in enumerate(GUIDE, start=1):
        ws.cell(i, 1, k)
        if v:
            ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(i, 1).font = Font(bold=True)
            ws.cell(i, 1).alignment = Alignment(vertical="top")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor="DBEAFE")
    wb.active = 0


def fill_example(wb, sheets):
    for sheet in sheets:
        rows = EXAMPLE[sheet]
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        for r, row in enumerate(rows, start=2):
            assert len(row) == len(headers), f"{sheet}: {len(row)} giá trị, {len(headers)} cột"
            for c, v in enumerate(row, start=1):
                if v is not None:
                    ws.cell(r, c, v)


ROUNDS = [
    ("1-phong-ban.xlsx", ["Phòng ban"]),
    ("2-danh-muc.xlsx", ["Chức vụ", "Bậc lương", "Danh sách ca", "Loại nghỉ phép"]),
    ("3-nhan-vien.xlsx", ["Nhân viên"]),
    ("4-so-du-phep-phan-ca.xlsx", ["Số dư phép", "Phân công ca"]),
]


def build(src, sheets, out, example):
    wb = openpyxl.load_workbook(src)
    for name in list(wb.sheetnames):
        if name not in sheets:
            del wb[name]
    add_guide(wb)
    if example:
        fill_example(wb, sheets)
    wb.save(out)


def main(argv):
    if len(argv) != 3:
        sys.exit(__doc__)
    src, out_dir = argv[1], argv[2]
    names = openpyxl.load_workbook(src, read_only=True).sheetnames
    for name in EXAMPLE:
        if name not in names:
            sys.exit(f"Thiếu sheet '{name}' trong file xuất (xuất bằng tiếng Việt, đủ 8 loại)")
    os.makedirs(os.path.join(out_dir, "vi-du"), exist_ok=True)
    for fname, sheets in ROUNDS:
        build(src, sheets, os.path.join(out_dir, fname), False)
        build(src, sheets, os.path.join(out_dir, "vi-du", fname), True)


if __name__ == "__main__":
    main(sys.argv)
